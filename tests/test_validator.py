"""Tests for presentation validator."""

import pytest
from pathlib import Path
from pptx import Presentation
from src.validator import parse_expected_slide_count, validate_presentation


def test_parse_expected_slide_count():
    """Test parsing expected slide count from task content."""
    task_content = """
    ## Slide 1 — Overview
    Content here

    ## Slide 2 — Details
    More content

    ## Slide 3 — Conclusion
    Final content
    """

    count = parse_expected_slide_count(task_content)
    assert count == 3


def test_parse_expected_slide_count_max():
    """Test that parser returns max slide number."""
    task_content = """
    ## Slide 1 — First
    ## Slide 5 — Last
    ## Slide 3 — Middle
    """

    count = parse_expected_slide_count(task_content)
    assert count == 5


def test_parse_expected_slide_count_ignores_h3_only():
    """H3/H4 headings must not be counted as H2 slide headers."""
    task_content = """
### Section A
#### Subsection
"""
    assert parse_expected_slide_count(task_content) == 0


def test_parse_expected_slide_count_h2_with_h3_between():
    """Real ## lines count; ### between them do not."""
    task_content = """
## First slide
### Sub only
## Second slide
"""
    assert parse_expected_slide_count(task_content) == 2


def test_parse_expected_slide_count_h3_slide_marker_ignored():
    """### Slide N must not contribute to max slide number."""
    task_content = """
### Slide 99 — bogus
## Slide 1 — real
## Slide 2 — real
"""
    assert parse_expected_slide_count(task_content) == 2


def test_parse_expected_slide_count_title_plus_h3_only():
    """H1 title slide counts; ### alone does not add body slides."""
    task_content = """# Main Title
### Only H3
#### H4
"""
    assert parse_expected_slide_count(task_content) == 1


def test_validate_presentation_missing_file():
    """Test validation of non-existent file."""
    result = validate_presentation(Path("nonexistent.pptx"), "")

    assert result.valid is False
    assert "does not exist" in result.issues[0]


def test_validate_presentation_valid(tmp_path):
    """Test validation of valid presentation."""
    # Create a valid presentation
    pptx_path = tmp_path / "test.pptx"
    prs = Presentation()

    # Add 3 slides with titles
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        t = slide.shapes.title
        assert t is not None
        t.text = f"Slide {i + 1}"

    prs.save(str(pptx_path))

    # Validate
    task_content = "## Slide 1\n## Slide 2\n## Slide 3"
    result = validate_presentation(pptx_path, task_content)

    assert result.slide_count == 3
    assert result.expected_slide_count == 3
    assert result.has_titles is True
    assert result.is_valid is True


def test_validate_presentation_wrong_count(tmp_path):
    """Test validation with wrong slide count."""
    pptx_path = tmp_path / "test.pptx"
    prs = Presentation()

    # Add 2 slides
    for i in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        t = slide.shapes.title
        assert t is not None
        t.text = f"Slide {i + 1}"

    prs.save(str(pptx_path))

    # Validate against 5 expected slides
    task_content = "## Slide 1\n## Slide 2\n## Slide 3\n## Slide 4\n## Slide 5"
    result = validate_presentation(pptx_path, task_content)

    assert result.slide_count == 2
    assert result.expected_slide_count == 5
    assert result.valid is False
    assert any("mismatch" in issue.lower() for issue in result.issues)
