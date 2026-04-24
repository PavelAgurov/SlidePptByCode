"""Sample working python-pptx code for testing."""

from pathlib import Path
from typing import Any, cast

from pptx import Presentation


def create_test_presentation():
    """Create a simple test presentation."""
    prs = Presentation()

    # Slide 1: Title
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    t1 = slide1.shapes.title
    assert t1 is not None
    t1.text = "Test Presentation"
    cast(Any, slide1.placeholders[1]).text = "A simple test"

    # Slide 2: Content
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    t2 = slide2.shapes.title
    assert t2 is not None
    t2.text = "Key Points"
    body = slide2.placeholders[1]
    tf = cast(Any, body).text_frame
    tf.text = "Point 1"
    tf.add_paragraph().text = "Point 2"
    tf.add_paragraph().text = "Point 3"

    # Slide 3: Conclusion
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    t3 = slide3.shapes.title
    assert t3 is not None
    t3.text = "Thank You"
    cast(Any, slide3.placeholders[1]).text = "Final remarks and contact information."

    # Save
    output_path = Path('.output') / 'test_fixture.pptx'
    output_path.parent.mkdir(exist_ok=True)
    prs.save(str(output_path))
    print(f"Presentation saved to {output_path}")


if __name__ == '__main__':
    create_test_presentation()
