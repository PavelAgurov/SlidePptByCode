"""Tests for ppt_bootstrap (init_deck, create_empty_ppt, strip_all_slides)."""

from pathlib import Path
from typing import Any, cast

from pptx import Presentation

from src.ppt_bootstrap import (
    copy_deck_template,
    create_empty_ppt,
    init_deck,
    strip_all_slides,
)
from src.task_chunker import TaskChunk


def test_create_empty_ppt_starts_with_zero_slides_and_has_layouts(
    tmp_path: Path,
) -> None:
    p = tmp_path / "empty.pptx"
    n = create_empty_ppt(p)
    assert p.exists()
    assert n == 0

    prs = Presentation(str(p))
    assert len(prs.slides) == 0
    assert len(prs.slide_layouts) >= 9


def test_create_empty_ppt_first_append_lands_at_index_zero(
    tmp_path: Path,
) -> None:
    p = tmp_path / "empty.pptx"
    create_empty_ppt(p)

    prs = Presentation(str(p))
    s = prs.slides.add_slide(prs.slide_layouts[1])
    st = s.shapes.title
    assert st is not None
    st.text = "Scratch slide"
    if len(s.placeholders) > 1:
        cast(Any, s.placeholders[1]).text = "Body text"
    prs.save(str(p))

    from src.validator import validate_scratch_append

    chunk = TaskChunk(kind="h2", title="T", body="## T\nx", index=1)
    v = validate_scratch_append(p, 0, chunk)
    assert v.is_valid
    assert v.slide_count == 1


def test_strip_all_slides_keeps_layouts(tmp_path: Path) -> None:
    p = tmp_path / "tpl.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.slides.add_slide(prs.slide_layouts[2])
    prs.save(str(p))
    assert len(Presentation(str(p)).slides) == 3

    n = strip_all_slides(p)
    assert n == 0
    prs2 = Presentation(str(p))
    assert len(prs2.slides) == 0
    assert len(prs2.slide_layouts) >= 9


def test_init_deck_without_template(tmp_path: Path) -> None:
    deck = tmp_path / "out" / "deck.pptx"
    n = init_deck(deck, None)
    assert n == 0
    assert deck.is_file()


def test_init_deck_with_template_strips_existing_slides(tmp_path: Path) -> None:
    tpl = tmp_path / "tpl.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(str(tpl))
    src_bytes = tpl.read_bytes()

    deck = tmp_path / "out" / "deck.pptx"
    n = init_deck(deck, tpl)
    assert n == 0
    prs2 = Presentation(str(deck))
    assert len(prs2.slides) == 0
    assert len(prs2.slide_layouts) >= 1

    # Template file untouched
    assert tpl.read_bytes() == src_bytes


def test_copy_deck_template_preserves_source(tmp_path: Path) -> None:
    src = tmp_path / "source.pptx"
    dst = tmp_path / "out" / "copy.pptx"
    create_empty_ppt(src)
    src_mtime = src.stat().st_mtime_ns
    src_bytes = src.read_bytes()

    n_dst = copy_deck_template(src, dst)
    assert dst.exists()
    assert n_dst == 0
    assert len(Presentation(str(dst)).slides) == 0

    assert src.read_bytes() == src_bytes
    assert src.stat().st_mtime_ns == src_mtime
