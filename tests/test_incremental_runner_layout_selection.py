from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest

from src.incremental_runner import run_incremental_pipeline
from src.layout_catalog import LayoutInfo
from src.models import LayoutDescription, LayoutSelection
from src.shared_default import DEFAULT_SHARED_PY


@pytest.fixture
def mock_config(tmp_path: Path) -> SimpleNamespace:
    out_dir = tmp_path / ".output"
    gen_dir = tmp_path / ".generated"
    out_dir.mkdir()
    gen_dir.mkdir()
    return SimpleNamespace(
        output_dir=out_dir,
        generated_dir=gen_dir,
    )


class _Executor:
    def __init__(self) -> None:
        self._saved: list[str] = []

    def save_code(self, _script: str, filename: str) -> Path:
        self._saved.append(filename)
        return Path(filename)

    def execute(
        self, _code_file: Path, *, expected_output_pptx: Path
    ) -> SimpleNamespace:
        return SimpleNamespace(
            success=True, error_message=None, stdout="", stderr="", traceback=""
        )


def _describe_layout_by_index(li: LayoutInfo) -> LayoutDescription:
    """Default: layout 0 = header; all other indices = content (valid for H2 filter)."""
    st = "header" if li.index == 0 else "content"
    return LayoutDescription(
        slide_type=st,  # type: ignore[arg-type]
        slide_has_image_placeholder=False,
        description="d",
    )


def _make_generator(*, select_first_offered: bool = True) -> Mock:
    g = Mock()
    g.extend_shared_module.return_value = DEFAULT_SHARED_PY
    g.generate_incremental_slide.return_value = SimpleNamespace(
        code="from pptx import Presentation\n"
    )
    g.fix_code_after_error.return_value = SimpleNamespace(
        code="from pptx import Presentation\n"
    )
    g.fix_incremental_h2_validation.return_value = SimpleNamespace(
        code="from pptx import Presentation\n"
    )
    g.describe_layout.side_effect = _describe_layout_by_index
    if select_first_offered:
        g.select_layout.side_effect = lambda *, layouts, **_k: LayoutSelection(
            explanation="x", selected_layout_index=layouts[0].index
        )
    return g


def _patch_validators(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "src.incremental_runner.build_shared_index",
        lambda *_a, **_k: "X=1",
    )
    monkeypatch.setattr(
        "src.incremental_runner.validate_scratch_append",
        lambda *a, **k: SimpleNamespace(is_valid=True, issues=[]),
    )
    monkeypatch.setattr(
        "src.incremental_runner.validate_deck_after_append",
        lambda *a, **k: SimpleNamespace(is_valid=True, issues=[]),
    )
    monkeypatch.setattr(
        "src.incremental_runner.validate_presentation",
        lambda *a, **k: SimpleNamespace(is_valid=True, issues=[]),
    )
    monkeypatch.setattr(
        "src.incremental_runner._sha256_file", lambda *_a, **_k: "0" * 64
    )


def test_layout_selection_called_for_every_slide_with_template(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    tpl = tmp_path / "tpl.pptx"
    from pptx import Presentation

    Presentation().save(str(tpl))

    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=0, name="Title", placeholders=()),
            LayoutInfo(index=1, name="Body", placeholders=()),
        ],
    )

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=tpl,
    )
    assert ok is True
    # 1 preamble + 1 H2 = 2 layout selections (all slides go through selection)
    assert generator.select_layout.call_count == 2


def test_layout_selection_called_even_without_template(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
    )
    assert ok is True
    # Default python-pptx layouts are still used; selection runs for every slide.
    assert generator.select_layout.call_count == 2


def test_forced_layout_h2_skips_h2_selection_only(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=0, name="Title Slide", placeholders=()),
            LayoutInfo(index=1, name="Content", placeholders=()),
        ],
    )

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
        forced_layout_h2="Content",
    )
    assert ok is True
    # Only preamble selection runs; H2 forced.
    assert generator.select_layout.call_count == 1


def test_forced_layout_both_skips_all_selection(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=0, name="Title Slide", placeholders=()),
            LayoutInfo(index=1, name="Content", placeholders=()),
        ],
    )

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
        forced_layout_h1="Title Slide",
        forced_layout_h2="Content",
    )
    assert ok is True
    assert generator.select_layout.call_count == 0


def test_forced_layout_unknown_name_errors(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [LayoutInfo(index=0, name="L0", placeholders=())],
    )

    ok, _deck, errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
        forced_layout_h2="NOPE",
    )
    assert ok is False
    assert any("layout name not found" in e for e in errs)


def test_preamble_layout_selection_failure_aborts(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    """If LLM returns out-of-range index twice for preamble, the pipeline aborts."""
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator(select_first_offered=False)
    # Both attempts return invalid index -> _select_layout returns None.
    generator.select_layout.return_value = LayoutSelection(
        explanation="oops", selected_layout_index=99
    )
    executor = _Executor()
    _patch_validators(monkeypatch)
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [LayoutInfo(index=0, name="L0", placeholders=())],
    )

    ok, _deck, errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
    )
    assert ok is False
    assert any("Preamble layout selection failed" in e for e in errs)


def test_h1_uses_only_header_layouts_in_selection(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)

    def _describe(li: LayoutInfo) -> LayoutDescription:
        if li.index == 0:
            return LayoutDescription(
                slide_type="content",
                slide_has_image_placeholder=False,
                description="body",
            )
        return LayoutDescription(
            slide_type="header",
            slide_has_image_placeholder=False,
            description="title",
        )

    generator.describe_layout.side_effect = _describe
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=0, name="Body", placeholders=()),
            LayoutInfo(index=1, name="Title", placeholders=()),
        ],
    )

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
    )
    assert ok is True
    preamble_call = generator.select_layout.call_args_list[0]
    offered = preamble_call.kwargs["layouts"]
    assert len(offered) == 1
    assert offered[0].index == 1
    assert offered[0].name == "Title"


def test_h2_excludes_header_and_other(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)

    types_by_index = {
        0: "header",
        1: "agenda",
        2: "content",
        3: "other",
    }

    def _describe(li: LayoutInfo) -> LayoutDescription:
        st = types_by_index[li.index]
        return LayoutDescription(
            slide_type=st,  # type: ignore[arg-type]
            slide_has_image_placeholder=False,
            description=f"idx{li.index}",
        )

    generator.describe_layout.side_effect = _describe
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=i, name=f"L{i}", placeholders=())
            for i in range(4)
        ],
    )

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
    )
    assert ok is True
    h2_call = generator.select_layout.call_args_list[1]
    offered = h2_call.kwargs["layouts"]
    indices = {li.index for li in offered}
    assert indices == {2}


def test_h2_excludes_layouts_with_image_placeholder(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)

    def _describe(li: LayoutInfo) -> LayoutDescription:
        if li.index == 0:
            return LayoutDescription(
                slide_type="header",
                slide_has_image_placeholder=False,
                description="h",
            )
        if li.index == 1:
            return LayoutDescription(
                slide_type="content",
                slide_has_image_placeholder=True,
                description="picture slot",
            )
        return LayoutDescription(
            slide_type="content",
            slide_has_image_placeholder=False,
            description="text only",
        )

    generator.describe_layout.side_effect = _describe
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=0, name="H", placeholders=()),
            LayoutInfo(index=1, name="Pic", placeholders=()),
            LayoutInfo(index=2, name="Body", placeholders=()),
        ],
    )

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
    )
    assert ok is True
    h2_call = generator.select_layout.call_args_list[1]
    offered = h2_call.kwargs["layouts"]
    assert {li.index for li in offered} == {2}


def test_other_layout_never_offered(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)

    def _describe(li: LayoutInfo) -> LayoutDescription:
        if li.index == 0:
            return LayoutDescription(
                slide_type="header",
                slide_has_image_placeholder=False,
                description="h1",
            )
        if li.index == 3:
            return LayoutDescription(
                slide_type="content",
                slide_has_image_placeholder=False,
                description="h2 body",
            )
        return LayoutDescription(
            slide_type="other",
            slide_has_image_placeholder=False,
            description="skip",
        )

    generator.describe_layout.side_effect = _describe
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=0, name="H0", placeholders=()),
            LayoutInfo(index=1, name="O1", placeholders=()),
            LayoutInfo(index=2, name="O2", placeholders=()),
            LayoutInfo(index=3, name="C3", placeholders=()),
        ],
    )

    ok, _deck, _errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
    )
    assert ok is True
    for call in generator.select_layout.call_args_list:
        offered = call.kwargs["layouts"]
        for li in offered:
            assert li.index != 1
            assert li.index != 2


def test_empty_filter_for_preamble_aborts(
    monkeypatch: pytest.MonkeyPatch, mock_config: SimpleNamespace, tmp_path: Path
) -> None:
    task = "# Title\n\n## Slide 1\n- A\n"
    generator = _make_generator()
    executor = _Executor()
    _patch_validators(monkeypatch)

    def _describe(_li: LayoutInfo) -> LayoutDescription:
        return LayoutDescription(
            slide_type="content",
            slide_has_image_placeholder=False,
            description="no headers",
        )

    generator.describe_layout.side_effect = _describe
    monkeypatch.setattr(
        "src.incremental_runner.read_layouts",
        lambda *_a, **_k: [
            LayoutInfo(index=0, name="A", placeholders=()),
            LayoutInfo(index=1, name="B", placeholders=()),
        ],
    )

    ok, _deck, errs = run_incremental_pipeline(
        task_content=task,
        style_content=None,
        language=None,
        output_filename=str(tmp_path / "out.pptx"),
        generator=generator,
        executor=executor,  # type: ignore[arg-type]
        config=mock_config,  # type: ignore[arg-type]
        max_retries=1,
        slide_max=None,
        template_pptx=None,
    )
    assert ok is False
    assert any("Preamble layout selection failed" in e for e in errs)
