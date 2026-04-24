"""Tests for incremental validation helpers."""

from pathlib import Path

from pptx import Presentation

from src.ppt_bootstrap import create_empty_ppt
from src.task_chunker import TaskChunk
from src.validator import (
    prefix_slide_digests,
    validate_deck_after_append,
    validate_incremental_h1_deck,
)


def test_prefix_slide_digests_stable(tmp_path: Path) -> None:
    p = tmp_path / "d.pptx"
    create_empty_ppt(p)
    prs = Presentation(str(p))
    s = prs.slides.add_slide(prs.slide_layouts[1])
    st = s.shapes.title
    assert st is not None
    st.text = "A"
    prs.save(str(p))

    d0 = prefix_slide_digests(p, 2)
    d1 = prefix_slide_digests(p, 2)
    assert d0 == d1


def test_validate_deck_after_append_prefix(tmp_path: Path) -> None:
    p = tmp_path / "deck.pptx"
    prs = Presentation()
    s0 = prs.slides.add_slide(prs.slide_layouts[1])
    t0 = s0.shapes.title
    assert t0 is not None
    t0.text = "First"
    prs.save(str(p))
    n_before = len(Presentation(str(p)).slides)

    prefix = prefix_slide_digests(p, n_before)
    prs2 = Presentation(str(p))
    s1 = prs2.slides.add_slide(prs2.slide_layouts[1])
    t1 = s1.shapes.title
    assert t1 is not None
    t1.text = "Second"
    prs2.save(str(p))

    chunk = TaskChunk(kind="h2", title="Second", body="## Second\nx", index=1)
    v = validate_deck_after_append(p, n_before, prefix, chunk)
    assert v.is_valid
    assert v.slide_count == n_before + 1


def test_validate_incremental_h1_deck_requires_shared(tmp_path: Path) -> None:
    deck = tmp_path / "d.pptx"
    create_empty_ppt(deck)
    prs = Presentation(str(deck))
    t = prs.slides[0].shapes.title
    assert t is not None
    t.text = "Hello Deck"
    prs.save(str(deck))

    shared = tmp_path / "shared.py"
    shared.write_text("# helpers\n", encoding="utf-8")

    v = validate_incremental_h1_deck(deck, shared, "Hello Deck")
    assert v.is_valid
