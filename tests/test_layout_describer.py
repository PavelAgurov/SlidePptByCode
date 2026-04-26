"""Tests for layout description cache and ensure_layout_descriptions."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import Mock

import pytest
from pptx import Presentation

from src.layout_catalog import LayoutInfo
from src.layout_describer import (
    cache_path,
    ensure_layout_descriptions,
    load_cached,
    save_cache,
    template_sha256,
)
from src.models import LayoutDescription
from src.prompts import LAYOUT_DESC_CACHE_PROMPT_VERSION


def test_cache_path_uses_stem_and_layout_subdir(tmp_path: Path) -> None:
    p = cache_path(tmp_path / "my" / "template_deck.pptx", tmp_path / ".generated")
    assert p == tmp_path / ".generated" / "layout" / "template_deck.json"


def test_template_sha256_deterministic(tmp_path: Path) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    a = template_sha256(t)
    b = template_sha256(t)
    assert a == b
    assert len(a) == 64


def test_load_cached_empty_when_missing_or_bad_sha(tmp_path: Path) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    sha = template_sha256(t)
    c = tmp_path / ".generated" / "layout" / "a.json"
    c.parent.mkdir(parents=True, exist_ok=True)
    c.write_text("not json", encoding="utf-8")
    assert load_cached(c, sha) == {}
    c.write_text(
        json.dumps(
            {
                "template_sha256": "0" * 64,
                "prompt_version": LAYOUT_DESC_CACHE_PROMPT_VERSION,
                "layouts": [{"index": 0, "name": "L", "description": "d"}],
            }
        ),
        encoding="utf-8",
    )
    assert load_cached(c, sha) == {}


def test_load_cached_rejects_wrong_prompt_version(tmp_path: Path) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    sha = template_sha256(t)
    c = tmp_path / ".generated" / "layout" / "a.json"
    c.parent.mkdir(parents=True, exist_ok=True)
    c.write_text(
        json.dumps(
            {
                "template_sha256": sha,
                "prompt_version": 9999,
                "layouts": [],
            }
        ),
        encoding="utf-8",
    )
    assert load_cached(c, sha) == {}


def test_load_cached_ok_partial_indices(tmp_path: Path) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    sha = template_sha256(t)
    c = tmp_path / ".generated" / "layout" / "a.json"
    c.parent.mkdir(parents=True, exist_ok=True)
    c.write_text(
        json.dumps(
            {
                "template_sha256": sha,
                "prompt_version": LAYOUT_DESC_CACHE_PROMPT_VERSION,
                "layouts": [
                    {"index": 0, "name": "A", "description": "one"},
                    {"index": 1, "name": "B", "description": "two"},
                ],
            }
        ),
        encoding="utf-8",
    )
    assert load_cached(c, sha) == {0: "one", 1: "two"}


def test_save_and_roundtrip(tmp_path: Path) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    sha = template_sha256(t)
    c = cache_path(t, tmp_path / ".gen")
    layouts = [
        LayoutInfo(index=0, name="L0", placeholders=()),
        LayoutInfo(index=1, name="L1", placeholders=()),
    ]
    save_cache(
        c,
        template_pptx=t,
        sha=sha,
        layouts=layouts,
        descriptions={0: "first", 1: "second"},
    )
    assert c.is_file()
    assert load_cached(c, sha) == {0: "first", 1: "second"}


def test_ensure_layout_descriptions_caches_per_layout_calls(
    tmp_path: Path,
) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    gen_dir = tmp_path / ".generated"
    gen_dir.mkdir()
    layouts = [
        LayoutInfo(index=0, name="A", placeholders=()),
        LayoutInfo(index=1, name="B", placeholders=()),
    ]
    g = Mock()
    g.describe_layout.side_effect = [
        LayoutDescription(description="d0"),
        LayoutDescription(description="d1"),
    ]
    d1 = ensure_layout_descriptions(
        generator=g, template_pptx=t, layouts=layouts, generated_dir=gen_dir
    )
    assert d1 == {0: "d0", 1: "d1"}
    assert g.describe_layout.call_count == 2
    g.reset_mock()
    d2 = ensure_layout_descriptions(
        generator=g, template_pptx=t, layouts=layouts, generated_dir=gen_dir
    )
    assert d2 == {0: "d0", 1: "d1"}
    assert g.describe_layout.call_count == 0


def test_ensure_layout_descriptions_fills_only_missing(
    tmp_path: Path,
) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    gen_dir = tmp_path / ".generated"
    gen_dir.mkdir()
    c = cache_path(t, gen_dir)
    c.parent.mkdir(parents=True, exist_ok=True)
    sha = template_sha256(t)
    c.write_text(
        json.dumps(
            {
                "template_sha256": sha,
                "prompt_version": LAYOUT_DESC_CACHE_PROMPT_VERSION,
                "layouts": [
                    {"index": 0, "name": "A", "description": "cached0"},
                ],
            }
        ),
        encoding="utf-8",
    )
    layouts = [
        LayoutInfo(index=0, name="A", placeholders=()),
        LayoutInfo(index=1, name="B", placeholders=()),
    ]
    g = Mock()
    g.describe_layout.return_value = LayoutDescription(description="new1")
    d = ensure_layout_descriptions(
        generator=g, template_pptx=t, layouts=layouts, generated_dir=gen_dir
    )
    assert d[0] == "cached0"
    assert d[1] == "new1"
    assert g.describe_layout.call_count == 1


def test_ensure_layout_descriptions_continues_on_llm_failure(
    tmp_path: Path,
) -> None:
    t = tmp_path / "a.pptx"
    Presentation().save(str(t))
    gen_dir = tmp_path / ".generated"
    gen_dir.mkdir()
    layouts = [
        LayoutInfo(index=0, name="A", placeholders=()),
        LayoutInfo(index=1, name="B", placeholders=()),
    ]
    g = Mock()
    g.describe_layout.side_effect = [LayoutDescription(description="ok0"), RuntimeError("boom")]

    d = ensure_layout_descriptions(
        generator=g, template_pptx=t, layouts=layouts, generated_dir=gen_dir
    )
    assert d == {0: "ok0"}
    assert g.describe_layout.call_count == 2
