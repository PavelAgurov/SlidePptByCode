"""Canonical python-pptx code snippets for LLM tool ``get_code_snippet``."""

from __future__ import annotations

import json
import logging
from enum import StrEnum
from typing import Any

logger = logging.getLogger(__name__)


class SnippetId(StrEnum):
    """Stable ids for ``get_code_snippet`` tool calls."""

    TITLE_SLIDE = "title_slide"
    CONTENT_SLIDE = "content_slide"
    BULLET_POINTS = "bullet_points"
    HEX_TO_RGB = "hex_to_rgb"
    BACKGROUND_COLOR = "background_color"
    TEXT_FORMATTING = "text_formatting"
    COLORED_SHAPE = "colored_shape"
    CHART_IMPORTS = "chart_imports"
    PARSE_CSV = "parse_csv"
    COLUMN_CHART = "column_chart"
    LINE_CHART = "line_chart"
    PIE_CHART = "pie_chart"
    STYLE_CHART_SERIES = "style_chart_series"
    TABLE_MARKDOWN_GRID = "table_markdown_grid"
    INCREMENTAL_H1_SKELETON = "incremental_h1_skeleton"
    INCREMENTAL_H2_SKELETON = "incremental_h2_skeleton"


SNIPPET_DESCRIPTIONS: dict[SnippetId, str] = {
    SnippetId.TITLE_SLIDE: "Add a title slide using layout 0 (title + subtitle placeholders).",
    SnippetId.CONTENT_SLIDE: (
        "Add a title-and-content slide: pick a layout with title + body placeholders "
        "via ``PP_PLACEHOLDER_TYPE`` (do not assume ``slide_layouts[1]`` or ``placeholders[1]``)."
    ),
    SnippetId.BULLET_POINTS: "Fill a body placeholder with bullets via TextFrame paragraphs.",
    SnippetId.HEX_TO_RGB: "Helper to parse #RRGGBB into RGB tuple for RGBColor.",
    SnippetId.BACKGROUND_COLOR: "Solid fill background on a slide.",
    SnippetId.TEXT_FORMATTING: "Font size, bold, and RGBColor on title/body paragraphs.",
    SnippetId.COLORED_SHAPE: "Add a filled rectangle (MSO_SHAPE) with line color.",
    SnippetId.CHART_IMPORTS: "Imports needed for CategoryChartData and chart enums.",
    SnippetId.PARSE_CSV: "Parse simple comma-separated rows from embedded task text.",
    SnippetId.COLUMN_CHART: "Blank layout + clustered column chart + textbox title.",
    SnippetId.LINE_CHART: "Line chart with markers and legend styling.",
    SnippetId.PIE_CHART: "Pie chart with legend on the right.",
    SnippetId.STYLE_CHART_SERIES: "Apply RGB fills to each series (uses hex_to_rgb helper).",
    SnippetId.TABLE_MARKDOWN_GRID: (
        "Parse GitHub-style pipe tables from markdown text and fill a native "
        "``slide.shapes.add_table`` (not a bullet list of raw '|' characters)."
    ),
    SnippetId.INCREMENTAL_H1_SKELETON: "H1 incremental pattern: write shared.py then open/save DECK_PPTX_PATH.",
    SnippetId.INCREMENTAL_H2_SKELETON: (
        "H2 incremental pattern: open TARGET_PPTX, append one slide using a layout with "
        "title + body placeholders by ``PP_PLACEHOLDER_TYPE``, save (portable across "
        "default scratch decks and corporate templates)."
    ),
}


SNIPPETS: dict[SnippetId, str] = {
    SnippetId.TITLE_SLIDE: """
slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title slide layout
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Main Title"
subtitle.text = "Subtitle"
""".strip(),
    SnippetId.CONTENT_SLIDE: """
from pptx.enum.shapes import PP_PLACEHOLDER_TYPE

_TITLE_TYPES = (
    PP_PLACEHOLDER_TYPE.TITLE,
    PP_PLACEHOLDER_TYPE.CENTER_TITLE,
    PP_PLACEHOLDER_TYPE.VERTICAL_TITLE,
)
_BODY_TYPES = (
    PP_PLACEHOLDER_TYPE.BODY,
    PP_PLACEHOLDER_TYPE.OBJECT,
    PP_PLACEHOLDER_TYPE.VERTICAL_BODY,
    PP_PLACEHOLDER_TYPE.VERTICAL_OBJECT,
)


def layout_has_title_and_body(layout) -> bool:
    has_title = False
    has_body = False
    for ph in layout.placeholders:
        t = ph.placeholder_format.type
        if t in _TITLE_TYPES:
            has_title = True
        if t in _BODY_TYPES:
            has_body = True
    return has_title and has_body


def pick_title_and_content_layout(prs):
    for layout in prs.slide_layouts:
        if layout_has_title_and_body(layout):
            return layout
    return prs.slide_layouts[1]


def first_body_placeholder(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.type in _BODY_TYPES:
            return ph
    raise ValueError("No body-like placeholder on slide")


slide_layout = pick_title_and_content_layout(prs)
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = first_body_placeholder(slide)
title.text = "Slide Title"
body.text = "Content goes here"
""".strip(),
    SnippetId.BULLET_POINTS: """
tf = body.text_frame
tf.text = "First bullet"
p = tf.add_paragraph()
p.text = "Second bullet"
p.level = 0  # Indentation level
""".strip(),
    SnippetId.HEX_TO_RGB: """
def hex_to_rgb(hex_color: str) -> tuple:
    # Accepts '#RRGGBB' or 'RRGGBB'; returns (r, g, b) ints 0-255
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))

# Usage:
r, g, b = hex_to_rgb('#76CCBE')
color = RGBColor(r, g, b)
""".strip(),
    SnippetId.BACKGROUND_COLOR: """
from pptx.dml.color import RGBColor

background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(118, 204, 190)
""".strip(),
    SnippetId.TEXT_FORMATTING: """
# Title paragraph
title = slide.shapes.title
title.text = "My Title"
title.text_frame.paragraphs[0].font.color.rgb = RGBColor(24, 139, 120)
title.text_frame.paragraphs[0].font.size = Pt(44)
title.text_frame.paragraphs[0].font.bold = True

# Body paragraphs
for paragraph in body.text_frame.paragraphs:
    paragraph.font.color.rgb = RGBColor(10, 33, 61)
    paragraph.font.size = Pt(18)
""".strip(),
    SnippetId.COLORED_SHAPE: """
from pptx.enum.shapes import MSO_SHAPE

left = Inches(1)
top = Inches(2)
width = Inches(3)
height = Inches(1)
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(250, 154, 38)
shape.line.color.rgb = RGBColor(24, 139, 120)
""".strip(),
    SnippetId.TABLE_MARKDOWN_GRID: '''
from pptx.util import Inches, Pt


def parse_pipe_markdown_table(markdown: str) -> list[list[str]]:
    """Extract rows from a GitHub-style pipe table embedded in markdown."""
    lines = [ln.rstrip() for ln in markdown.splitlines()]
    pipe_lines = [
        ln for ln in lines if ln.strip().startswith("|") and ln.count("|") >= 2
    ]
    if not pipe_lines:
        return []

    def is_separator_row(line: str) -> bool:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if not cells:
            return False
        for c in cells:
            if not c or not all(ch in "-: " for ch in c):
                return False
        return True

    grid: list[list[str]] = []
    for ln in pipe_lines:
        if is_separator_row(ln):
            continue
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        grid.append(cells)
    return grid


def fill_table_from_grid(slide, grid: list[list[str]], left, top, width, height) -> None:
    """Create a real PowerPoint table and copy cell text (not one bullet of pipe text)."""
    if not grid:
        return
    nrows = len(grid)
    ncols = max(len(r) for r in grid)
    tbl_shape = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    tbl = tbl_shape.table
    for r in range(nrows):
        for c in range(ncols):
            txt = grid[r][c] if c < len(grid[r]) else ""
            cell = tbl.cell(r, c)
            cell.text = txt
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(8)

# Typical call on a content slide (hide or clear the body placeholder if you only show the table):
# grid = parse_pipe_markdown_table(section_markdown)
# fill_table_from_grid(slide, grid, Inches(0.35), Inches(1.35), Inches(9.3), Inches(5.0))
'''.strip(),
    SnippetId.CHART_IMPORTS: """
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
""".strip(),
    SnippetId.PARSE_CSV: '''
# Example: parse CSV text embedded in the task
csv_text = """month,product_A,product_B,product_C
2025-01,120,95,60
2025-02,135,102,72
2025-03,150,110,80"""

lines = [line.strip() for line in csv_text.strip().split("\n") if line.strip()]
headers = lines[0].split(",")
categories = []
series_data = {header: [] for header in headers[1:]}

for line in lines[1:]:
    values = line.split(",")
    categories.append(values[0])
    for i, header in enumerate(headers[1:], 1):
        series_data[header].append(float(values[i]))
'''.strip(),
    SnippetId.COLUMN_CHART: """
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

chart_data = CategoryChartData()
chart_data.categories = categories
for series_name, values in series_data.items():
    chart_data.add_series(series_name, values)

x, y, cx, cy = Inches(1), Inches(1.5), Inches(8), Inches(5)
chart = slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
).chart

title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(0.8))
title_frame = title_box.text_frame
title_frame.text = "Sales Data"
title_frame.paragraphs[0].font.size = Pt(32)
title_frame.paragraphs[0].font.bold = True
""".strip(),
    SnippetId.LINE_CHART: """
slide = prs.slides.add_slide(prs.slide_layouts[6])

chart_data = CategoryChartData()
chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
chart_data.add_series("Revenue", [120, 135, 150, 170])
chart_data.add_series("Costs", [95, 102, 110, 130])

x, y, cx, cy = Inches(1), Inches(1.5), Inches(8), Inches(5)
chart = slide.shapes.add_chart(
    XL_CHART_TYPE.LINE_MARKERS, x, y, cx, cy, chart_data
).chart

chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.BOTTOM
chart.legend.font.size = Pt(12)
""".strip(),
    SnippetId.PIE_CHART: """
slide = prs.slides.add_slide(prs.slide_layouts[6])

chart_data = CategoryChartData()
chart_data.categories = ["Product A", "Product B", "Product C"]
chart_data.add_series("Market Share", [45, 30, 25])

x, y, cx, cy = Inches(2), Inches(1.5), Inches(6), Inches(5)
chart = slide.shapes.add_chart(XL_CHART_TYPE.PIE, x, y, cx, cy, chart_data).chart

chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.RIGHT
chart.legend.font.size = Pt(11)
""".strip(),
    SnippetId.STYLE_CHART_SERIES: """
# Requires hex_to_rgb defined (see hex_to_rgb snippet)
series_colors = [
    hex_to_rgb("#188B78"),
    hex_to_rgb("#FA9A26"),
    hex_to_rgb("#1998DD"),
    hex_to_rgb("#0A213D"),
]

for idx, series in enumerate(chart.series):
    if idx < len(series_colors):
        fill = series.format.fill
        fill.solid()
        r, g, b = series_colors[idx]
        fill.fore_color.rgb = RGBColor(r, g, b)

for series in chart.series:
    series.has_data_labels = True
    series.data_labels.font.size = Pt(10)
""".strip(),
    SnippetId.INCREMENTAL_H1_SKELETON: """
# Paths DECK_PPTX_PATH and SHARED_PY_PATH are injected by the orchestrator.
from pathlib import Path
from pptx import Presentation

SHARED_SRC = '''\
# Example shared module body (replace with your real palette + helpers)
from pptx.dml.color import RGBColor

ACCENT = RGBColor(24, 139, 120)
'''


def main() -> None:
    SHARED_PY_PATH.write_text(SHARED_SRC, encoding="utf-8")
    prs = Presentation(str(DECK_PPTX_PATH))
    # Edit existing title slide(s) or add minimal H1 slides here
    prs.save(str(DECK_PPTX_PATH))


if __name__ == "__main__":
    main()
""".strip(),
    SnippetId.INCREMENTAL_H2_SKELETON: """
# TARGET_PPTX and pre-loaded ``shared`` are injected by the orchestrator.
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER_TYPE

_TITLE_TYPES = (
    PP_PLACEHOLDER_TYPE.TITLE,
    PP_PLACEHOLDER_TYPE.CENTER_TITLE,
    PP_PLACEHOLDER_TYPE.VERTICAL_TITLE,
)
_BODY_TYPES = (
    PP_PLACEHOLDER_TYPE.BODY,
    PP_PLACEHOLDER_TYPE.OBJECT,
    PP_PLACEHOLDER_TYPE.VERTICAL_BODY,
    PP_PLACEHOLDER_TYPE.VERTICAL_OBJECT,
)


def layout_has_title_and_body(layout) -> bool:
    has_title = False
    has_body = False
    for ph in layout.placeholders:
        t = ph.placeholder_format.type
        if t in _TITLE_TYPES:
            has_title = True
        if t in _BODY_TYPES:
            has_body = True
    return has_title and has_body


def pick_title_and_content_layout(prs):
    for layout in prs.slide_layouts:
        if layout_has_title_and_body(layout):
            return layout
    return prs.slide_layouts[1]


def first_body_placeholder(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.type in _BODY_TYPES:
            return ph
    raise ValueError("No body-like placeholder on slide")


def main() -> None:
    prs = Presentation(str(TARGET_PPTX))
    layout = pick_title_and_content_layout(prs)
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Slide title"
    body = first_body_placeholder(slide)
    body.text = "Body text"
    prs.save(str(TARGET_PPTX))


if __name__ == "__main__":
    main()
""".strip(),
}


class SnippetCallCache:
    """Per-LLM-call dedup of snippet bodies (do not persist across slides)."""

    def __init__(self, *, max_reuse_hints: int = 4) -> None:
        self._issued: set[str] = set()
        self._last_sid: str | None = None
        self._reuse_streak = 0
        self._max_reuse_hints = max_reuse_hints

    def get_body(self, snippet_id: str) -> str:
        try:
            key = SnippetId(snippet_id)
        except ValueError:
            valid = ", ".join(s.value for s in SnippetId)
            return f"Unknown snippet_id={snippet_id!r}. Valid ids: {valid}"

        body = SNIPPETS[key]
        if snippet_id not in self._issued:
            self._issued.add(snippet_id)
            self._last_sid = snippet_id
            self._reuse_streak = 0
            logger.info(
                "snippet cache: first fetch id=%s chars=%d",
                snippet_id,
                len(body),
                extra={"color_event": "tool_result"},
            )
            return body

        if self._last_sid == snippet_id:
            self._reuse_streak += 1
        else:
            self._last_sid = snippet_id
            self._reuse_streak = 1

        if self._reuse_streak >= self._max_reuse_hints:
            raise RuntimeError(
                f"get_code_snippet({snippet_id!r}) repeated too many times after reuse hint"
            )

        logger.info(
            "snippet cache: reuse hint id=%s streak=%d",
            snippet_id,
            self._reuse_streak,
            extra={"color_event": "snippet_reuse"},
        )
        return (
            f"(Snippet {snippet_id!r} was already returned earlier in this turn — "
            "reuse that exact text from the conversation; do not request it again.)"
        )


def handle_get_code_snippet(arguments_json: str, cache: SnippetCallCache) -> str:
    """Parse tool arguments and return snippet text or error string."""
    try:
        args = json.loads(arguments_json or "{}")
    except json.JSONDecodeError as e:
        return f"Invalid JSON in tool arguments: {e}"
    sid = args.get("snippet_id", "")
    if not isinstance(sid, str):
        return "snippet_id must be a string"
    return cache.get_body(sid)


def get_code_snippet_tools() -> list[dict[str, Any]]:
    """OpenAI-compatible tool list for ``beta.chat.completions.parse``."""
    enum_vals = [s.value for s in SnippetId]
    return [
        {
            "type": "function",
            "function": {
                "name": "get_code_snippet",
                "strict": True,
                "description": (
                    "Return canonical python-pptx example code for the given id. "
                    "Call before writing non-trivial patterns; do not invent APIs."
                ),
                "parameters": {
                    "type": "object",
                    "properties": {
                        "snippet_id": {
                            "type": "string",
                            "enum": enum_vals,
                            "description": "Snippet identifier",
                        }
                    },
                    "required": ["snippet_id"],
                    "additionalProperties": False,
                },
            },
        }
    ]


def formatted_snippet_catalog_for_prompt() -> str:
    """Multi-line catalog for system prompts (ids only, no full code)."""
    lines = [
        "CANONICAL CODE SNIPPETS — call tool ``get_code_snippet`` with ``snippet_id``:",
        "",
    ]
    for sid in SnippetId:
        lines.append(f"- {sid.value}: {SNIPPET_DESCRIPTIONS[sid]}")
    return "\n".join(lines)
