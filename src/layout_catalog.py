"""Read slide layouts (masters) and placeholders from a .pptx template.

This module is used for selecting the best layout for a slide when a user
provides a template deck.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation

from src.models import LayoutDescription


@dataclass(frozen=True, slots=True)
class LayoutPlaceholder:
    idx: int
    type: str
    name: str
    prompt_text: str
    left_in: float
    top_in: float
    width_in: float
    height_in: float


@dataclass(frozen=True, slots=True)
class LayoutShape:
    """Non-placeholder shape on a slide layout (decorative / structural)."""

    name: str
    shape_type: str
    auto_shape_type: str
    fill_summary: str
    left_in: float
    top_in: float
    width_in: float
    height_in: float


@dataclass(frozen=True, slots=True)
class LayoutInfo:
    index: int
    name: str
    placeholders: tuple[LayoutPlaceholder, ...]
    decor_shapes: tuple[LayoutShape, ...] = ()


def _safe_inches(value: object) -> float:
    """
    Convert an EMU-like value to inches as float.

    python-pptx returns shape coordinates as Emu, which provides `.inches`.
    Some placeholders may have missing geometry; in that case return 0.0.
    """

    try:
        inches = float(getattr(value, "inches"))
    except Exception:
        return 0.0
    # Keep reasonably compact for prompt payloads.
    return round(inches, 2)


def _ph_type_name(ph) -> str:
    try:
        t = ph.placeholder_format.type
    except Exception:
        return "UNKNOWN"
    try:
        # python-pptx formats this as e.g. "TITLE (1)".
        raw = str(t).strip()
        head = raw.split(".", 1)[-1]
        head = head.split("(", 1)[0].strip()
        head = head.split(" ", 1)[0].strip()
        return head or "UNKNOWN"
    except Exception:
        return "UNKNOWN"


def _shape_type_str(shape) -> str:
    try:
        st = shape.shape_type
    except Exception:
        return "UNKNOWN"
    try:
        raw = str(st).strip()
        head = raw.split(".", 1)[-1]
        head = head.split("(", 1)[0].strip()
        head = head.split(" ", 1)[0].strip()
        return head or "UNKNOWN"
    except Exception:
        return "UNKNOWN"


def _auto_shape_type_str(shape) -> str:
    try:
        ast = getattr(shape, "auto_shape_type", None)
        if ast is None:
            return ""
        raw = str(ast).strip()
        head = raw.split(".", 1)[-1]
        head = head.split("(", 1)[0].strip()
        head = head.split(" ", 1)[0].strip()
        return head or ""
    except Exception:
        return ""


def _fill_summary(shape) -> str:
    """Compact fill hint for LLM (solid hex, none, picture, gradient, or empty)."""
    try:
        fill = getattr(shape, "fill", None)
        if fill is None:
            return ""
        try:
            ft = fill.type
        except Exception:
            return ""
        try:
            type_name = str(ft).split(".")[-1].split("(", 1)[0].strip().upper()
        except Exception:
            type_name = ""

        if "PICTURE" in type_name or type_name == "PICTURE":
            return "picture"
        if "GRADIENT" in type_name or type_name == "GRADIENT":
            return "gradient"
        if "BACKGROUND" in type_name:
            return "background"
        if type_name in ("NONE", ""):
            try:
                if getattr(fill, "fore_color", None) is not None:
                    pass
            except Exception:
                pass
            return "none"

        # MSO_FILL_TYPE.SOLID etc.
        if "SOLID" in type_name or type_name == "SOLID":
            try:
                rgb = fill.fore_color.rgb
                if rgb is not None:
                    return f"solid:#{rgb}"
            except Exception:
                pass
            return "solid"
        return type_name.lower() or ""
    except Exception:
        return ""


def _read_decor_shapes(layout) -> tuple[LayoutShape, ...]:
    out: list[LayoutShape] = []
    try:
        shapes_iter = layout.shapes
    except Exception:
        return ()
    for shape in shapes_iter:
        try:
            if shape.is_placeholder:
                continue
        except Exception:
            continue
        try:
            nm = str(getattr(shape, "name", "") or "")
        except Exception:
            nm = ""
        out.append(
            LayoutShape(
                name=nm,
                shape_type=_shape_type_str(shape),
                auto_shape_type=_auto_shape_type_str(shape),
                fill_summary=_fill_summary(shape),
                left_in=_safe_inches(getattr(shape, "left", None)),
                top_in=_safe_inches(getattr(shape, "top", None)),
                width_in=_safe_inches(getattr(shape, "width", None)),
                height_in=_safe_inches(getattr(shape, "height", None)),
            )
        )
    return tuple(out)


def read_layouts(pptx_path: Path) -> list[LayoutInfo]:
    """
    Read all slide layouts and their placeholders from a pptx deck.

    Raises on invalid pptx / IO errors; callers should decide whether to
    fallback to default behavior.
    """

    prs = Presentation(str(Path(pptx_path)))
    out: list[LayoutInfo] = []

    for i, layout in enumerate(prs.slide_layouts):
        placeholders: list[LayoutPlaceholder] = []
        for ph in layout.placeholders:
            try:
                idx = int(ph.placeholder_format.idx)
            except Exception:
                # If idx is not accessible we cannot address the placeholder reliably.
                continue

            try:
                name = str(getattr(ph, "name", "")) or ""
            except Exception:
                name = ""

            prompt_text = ""
            try:
                tf = getattr(ph, "text_frame", None)
                if tf is not None:
                    prompt_text = (tf.text or "").strip()
            except Exception:
                prompt_text = ""

            placeholders.append(
                LayoutPlaceholder(
                    idx=idx,
                    type=_ph_type_name(ph),
                    name=name,
                    prompt_text=prompt_text,
                    left_in=_safe_inches(getattr(ph, "left", None)),
                    top_in=_safe_inches(getattr(ph, "top", None)),
                    width_in=_safe_inches(getattr(ph, "width", None)),
                    height_in=_safe_inches(getattr(ph, "height", None)),
                )
            )

        decor = _read_decor_shapes(layout)

        out.append(
            LayoutInfo(
                index=i,
                name=str(getattr(layout, "name", "")) or "",
                placeholders=tuple(placeholders),
                decor_shapes=decor,
            )
        )

    return out


def format_layout_card_full(li: LayoutInfo) -> str:
    lines = [
        f'### Layout #{li.index} — \"{li.name}\"',
        "Placeholders:",
    ]
    if not li.placeholders:
        lines.append("- (no placeholders)")
    else:
        for ph in li.placeholders:
            lines.append(
                (
                    f'- idx={ph.idx} type={ph.type} name=\"{ph.name}\" '
                    f'prompt=\"{ph.prompt_text}\" '
                    f'bbox=({ph.left_in},{ph.top_in},{ph.width_in}x{ph.height_in}) in'
                )
            )
    lines.append("Decorative shapes:")
    if not li.decor_shapes:
        lines.append("- (none)")
    else:
        for ds in li.decor_shapes:
            ast = f' autoshape={ds.auto_shape_type!r}' if ds.auto_shape_type else ""
            fill = f' fill={ds.fill_summary!r}' if ds.fill_summary else ""
            lines.append(
                (
                    f'- name=\"{ds.name}\" type={ds.shape_type}{ast}{fill} '
                    f'bbox=({ds.left_in},{ds.top_in},{ds.width_in}x{ds.height_in}) in'
                )
            )
    return "\n".join(lines)


def format_layout_card_compact(li: LayoutInfo) -> str:
    types = ", ".join(ph.type for ph in li.placeholders) if li.placeholders else ""
    if not types:
        types = "(none)"
    decor_n = len(li.decor_shapes)
    decor_bit = f", decor_shapes: {decor_n}" if decor_n else ""
    return (
        f'- #{li.index} \"{li.name}\" — placeholders: {len(li.placeholders)} ({types}){decor_bit}'
    )


def format_layouts_catalog_compact(items: list[LayoutInfo]) -> str:
    if not items:
        return "(no layouts)"
    return "\n".join(format_layout_card_compact(x) for x in items)


def format_layouts_catalog_with_descriptions(
    items: list[LayoutInfo],
    descriptions: dict[int, LayoutDescription] | None,
) -> str:
    """
    Like ``format_layouts_catalog_compact`` but appends ``desc:`` and ``zones=N`` per layout
    when a matching ``LayoutDescription`` exists for that index.
    """

    if not items:
        return "(no layouts)"
    if not descriptions:
        return format_layouts_catalog_compact(items)
    out: list[str] = []
    for li in items:
        line = format_layout_card_compact(li)
        ld = descriptions.get(li.index)
        if ld is not None:
            d = (ld.description or "").strip()
            if d:
                line = f"{line}  desc: {d}"
            line = f"{line}  zones={ld.content_zones_count}"
        out.append(line)
    return "\n".join(out)
