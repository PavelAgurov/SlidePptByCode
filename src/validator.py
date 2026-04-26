"""Presentation validation logic."""

import logging
import re
from pathlib import Path
from pptx import Presentation
from .models import ValidationResult
from .task_chunker import TaskChunk

logger = logging.getLogger(__name__)


def slide_text_digest(slide) -> str:
    """Compact text fingerprint for a slide (title + text shapes)."""
    parts: list[str] = []
    try:
        if slide.shapes.title and slide.shapes.title.text:
            parts.append(slide.shapes.title.text.strip())
    except Exception:
        pass
    for shape in slide.shapes:
        try:
            if hasattr(shape, "text_frame") and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t[:240])
        except Exception:
            continue
    return (" | ".join(parts))[:800]


def prefix_slide_digests(pptx_path: Path, num_slides: int) -> list[str]:
    """Return ``slide_text_digest`` for slides ``0 .. num_slides-1``."""
    prs = Presentation(str(pptx_path))
    n = min(num_slides, len(prs.slides))
    return [slide_text_digest(prs.slides[i]) for i in range(n)]


def validate_scratch_append(
    scratch_path: Path,
    baseline_slide_count: int,
    _h2_chunk: TaskChunk,
) -> ValidationResult:
    """After slide script on scratch: slide count +1 and new slide has content."""
    issues: list[str] = []
    if not scratch_path.exists():
        return ValidationResult(
            valid=False,
            slide_count=0,
            has_titles=False,
            issues=["Scratch pptx does not exist"],
        )
    try:
        prs = Presentation(str(scratch_path))
        slide_count = len(prs.slides)
        expected = baseline_slide_count + 1
        if slide_count != expected:
            issues.append(
                f"Scratch slide count expected {expected} (baseline {baseline_slide_count}+1), "
                f"got {slide_count}"
            )
        if slide_count > 0:
            last = slide_text_digest(prs.slides[-1])
            if not last.strip():
                issues.append("Appended scratch slide has no readable text/content")
        valid = len(issues) == 0
        return ValidationResult(
            valid=valid,
            slide_count=slide_count,
            expected_slide_count=expected,
            has_titles=slide_count > 0,
            issues=issues,
        )
    except Exception as e:
        return ValidationResult(
            valid=False,
            slide_count=0,
            has_titles=False,
            issues=[f"Error validating scratch: {e}"],
        )


def validate_deck_after_append(
    deck_path: Path,
    previous_slide_count: int,
    prefix_digests_before: list[str],
    _h2_chunk: TaskChunk,
) -> ValidationResult:
    """
    After appending one slide to deck: count +1, prefix digests unchanged, new slide ok.
    """
    issues: list[str] = []
    if not deck_path.exists():
        return ValidationResult(
            valid=False,
            slide_count=0,
            has_titles=False,
            issues=["Deck file does not exist"],
        )
    try:
        prs = Presentation(str(deck_path))
        slide_count = len(prs.slides)
        expected = previous_slide_count + 1
        if slide_count != expected:
            issues.append(
                f"Deck slide count expected {expected} (was {previous_slide_count}+1), "
                f"got {slide_count}"
            )
        # Prefix: existing slides must be unchanged (digest match)
        if previous_slide_count > 0:
            after_prefix = prefix_slide_digests(deck_path, previous_slide_count)
            if after_prefix != prefix_digests_before:
                issues.append(
                    "Deck prefix slides changed (append-only violation or corrupted deck)"
                )
        if slide_count > 0:
            last = slide_text_digest(prs.slides[-1])
            if not last.strip():
                issues.append("New deck slide has no readable text/content")
        valid = len(issues) == 0
        return ValidationResult(
            valid=valid,
            slide_count=slide_count,
            expected_slide_count=expected,
            has_titles=slide_count > 0,
            issues=issues,
        )
    except Exception as e:
        return ValidationResult(
            valid=False,
            slide_count=0,
            has_titles=False,
            issues=[f"Error validating deck after append: {e}"],
        )


def parse_expected_slide_count(task_content: str) -> int:
    """
    Parse expected slide count from task content.

    Looks for patterns like "## Slide N" to count expected slides.
    Also checks for H1 header (# Title) which indicates a title slide.

    Args:
        task_content: Text content of the task

    Returns:
        Expected number of slides
    """
    # Check if there's a title slide (H1 header at the beginning)
    has_title_slide = bool(re.search(r'^\s*#\s+[^#\n]+', task_content, re.MULTILINE))

    # Find all "## Slide N" patterns (H2 only, not ### Slide N)
    slide_markers = re.findall(r'(?mi)^\s*#{2}(?!#)\s+Slide\s+(\d+)', task_content)

    if slide_markers:
        # Get the highest slide number
        max_slide = max(int(num) for num in slide_markers)
        # Add 1 if there's a title slide
        total_slides = max_slide + (1 if has_title_slide else 0)
        logger.debug(f"Found {len(slide_markers)} slide markers, max: {max_slide}, title slide: {has_title_slide}, total: {total_slides}")
        return total_slides

    # Fallback: count lines that look like H2 slide headers (not ### etc.)
    slide_headers = re.findall(r'(?m)^\s*#{2}(?!#)\s+[^\n]+', task_content)
    total_slides = len(slide_headers) + (1 if has_title_slide else 0)
    logger.debug(f"Fallback: found {len(slide_headers)} potential slide headers, title slide: {has_title_slide}, total: {total_slides}")
    return total_slides


def validate_presentation(
    pptx_path: Path,
    task_content: str
) -> ValidationResult:
    """
    Validate a generated presentation.

    Checks:
    - File exists and is not empty
    - Can be opened as a valid .pptx
    - Slide count matches expected
    - Slides have titles

    Args:
        pptx_path: Path to .pptx file
        task_content: Original task specification

    Returns:
        ValidationResult with validation details
    """
    logger.info(f"Validating presentation: {pptx_path}")

    issues = []

    # Check file exists
    if not pptx_path.exists():
        logger.error(f"File does not exist: {pptx_path}")
        return ValidationResult(
            valid=False,
            slide_count=0,
            has_titles=False,
            issues=["File does not exist"]
        )

    # Check file is not empty
    if pptx_path.stat().st_size == 0:
        logger.error(f"File is empty: {pptx_path}")
        return ValidationResult(
            valid=False,
            slide_count=0,
            has_titles=False,
            issues=["File is empty"]
        )

    try:
        # Open presentation
        prs = Presentation(str(pptx_path))
        slide_count = len(prs.slides)
        logger.info(f"Presentation has {slide_count} slides")

        # Parse expected slide count
        expected_slide_count = parse_expected_slide_count(task_content)
        logger.info(f"Expected {expected_slide_count} slides")

        # Check slide count
        if slide_count != expected_slide_count:
            issues.append(
                f"Slide count mismatch: expected {expected_slide_count}, got {slide_count}"
            )

        # Check if slides have titles or meaningful content (charts, text boxes)
        slides_with_content = 0
        for i, slide in enumerate(prs.slides, 1):
            has_title = False
            has_other_content = False

            # Check for title shape with text
            if slide.shapes.title:
                title_text = slide.shapes.title.text.strip()
                if title_text:
                    has_title = True
                    slides_with_content += 1

            # If no title, check for chart or text box content
            if not has_title:
                for shape in slide.shapes:
                    # Check for chart
                    if hasattr(shape, 'chart'):
                        has_other_content = True
                        break
                    # Check for text box with content
                    tf = getattr(shape, "text_frame", None)
                    if tf is not None:
                        text = tf.text.strip()
                        if text:
                            has_other_content = True
                            break

                if has_other_content:
                    slides_with_content += 1
                else:
                    issues.append(f"Slide {i} has no title or content")

        has_titles = slides_with_content == slide_count

        # Determine if valid
        valid = len(issues) == 0

        logger.info(f"Validation result: valid={valid}, issues={len(issues)}")

        return ValidationResult(
            valid=valid,
            slide_count=slide_count,
            expected_slide_count=expected_slide_count,
            has_titles=has_titles,
            issues=issues
        )

    except Exception as e:
        logger.error(f"Error validating presentation: {e}")
        return ValidationResult(
            valid=False,
            slide_count=0,
            has_titles=False,
            issues=[f"Error opening presentation: {str(e)}"]
        )
